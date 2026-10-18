#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Construye Meteo_ValledelEria_datos.xlsx a partir de los CSV de la carpeta datos/.

Hojas:
  - Leeme   : qué es cada cosa y qué avisos tener en cuenta.
  - Diario  : un resumen por día y estación (T mín / máx / media, lluvia del día,
              lluvia acumulada 7 y 30 días, horas con dato).
  - Horario : todos los datos hora a hora, lado a lado. Los huecos quedan VACÍOS.

Necesita: openpyxl.  Se ejecuta tras captura_saih.py.

El libro se genera entero cada vez, así que NO es para editar a mano. Para anotar
cosas (observaciones propias) se usa otro archivo, aparte.
Los valores se escriben ya calculados (no fórmulas) a propósito: los archivos que se
descargan de internet se abren en "Vista protegida" de Excel, que no calcula
fórmulas hasta pulsar "Habilitar edición"; así se ven bien desde el primer segundo.
"""
import csv
import json
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from captura_saih import CARPETA, SERIES, TZ, calidad, leer_csv

SALIDA = Path(os.environ.get("SALIDA_XLSX", "Meteo_ValledelEria_datos.xlsx"))

ESTACIONES = [
    ("castrocalbon", "Castrocalbón (PL284)"),
    ("moralesdelrey", "Morales del Rey (EA082)"),
]

FUENTE = "Arial"
F_NORMAL = Font(name=FUENTE, size=10)
F_NEGRITA = Font(name=FUENTE, size=10, bold=True)
F_TITULO = Font(name=FUENTE, size=14, bold=True)
F_CAB = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
REL_CAB = PatternFill("solid", fgColor="1F4E78")
REL_GRUPO = PatternFill("solid", fgColor="2E75B6")
REL_INCOMPLETO = PatternFill("solid", fgColor="F8CBAD")  # naranja claro
REL_DUDOSO = PatternFill("solid", fgColor="FFE699")      # amarillo
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def horas_del_dia(d):
    """Horas que tiene ese día local (24, o 23/25 en los cambios de hora)."""
    siguiente = d + timedelta(days=1)
    ini = datetime(d.year, d.month, d.day, tzinfo=TZ).astimezone(timezone.utc)
    fin = datetime(siguiente.year, siguiente.month, siguiente.day, tzinfo=TZ).astimezone(timezone.utc)
    return round((fin - ini).total_seconds() / 3600)


def cargar():
    """{serie: {epoch: (valor, calidad)}}"""
    series = {}
    for nombre, cfg in SERIES.items():
        crudo = leer_csv(CARPETA / f"{nombre}.csv")
        series[nombre] = {e: (v, calidad(cfg["tipo"], v)) for e, v in crudo.items()}
    return series


def agrupar_por_dia(serie):
    por_dia = defaultdict(list)
    for epoch, (valor, cal) in serie.items():
        por_dia[datetime.fromtimestamp(epoch, TZ).date()].append((valor, cal))
    return por_dia


def hoja_horario(wb, series):
    ws = wb.create_sheet("Horario")
    cabeceras = ["Fecha y hora (local)"]
    nombres = []
    for clave, etiqueta in ESTACIONES:
        cabeceras += [f"{etiqueta}\nTemperatura (ºC)", f"{etiqueta}\nLluvia (l/m²)"]
        nombres += [f"{clave}_temperatura", f"{clave}_lluvia"]
    for c, texto in enumerate(cabeceras, start=1):
        celda = ws.cell(row=1, column=c, value=texto)
        celda.font, celda.fill, celda.alignment = F_CAB, REL_CAB, CENTRO
    ws.row_dimensions[1].height = 42

    epochs = sorted(set().union(*[set(s) for s in series.values()]))
    for i, epoch in enumerate(epochs, start=2):
        momento = datetime.fromtimestamp(epoch, TZ).replace(tzinfo=None)
        c = ws.cell(row=i, column=1, value=momento)
        c.number_format, c.font = "dd/mm/yyyy hh:mm", F_NORMAL
        for j, nombre in enumerate(nombres, start=2):
            dato = series[nombre].get(epoch)
            celda = ws.cell(row=i, column=j)
            celda.font = F_NORMAL
            if dato is not None:
                celda.value = dato[0]
                celda.number_format = "0.0" if nombre.endswith("temperatura") else "0.0"
                if dato[1] == "dudoso":
                    celda.fill = REL_DUDOSO
    ws.column_dimensions["A"].width = 20
    for col in range(2, 6):
        ws.column_dimensions[get_column_letter(col)].width = 24
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:E{max(len(epochs) + 1, 2)}"
    return len(epochs)


PREFIJOS = {"castrocalbon": "cas", "moralesdelrey": "mor"}
SUB = ["T mín (ºC)", "T máx (ºC)", "T media (ºC)", "Horas con dato T",
       "Lluvia del día (l/m²)", "Horas con dato lluvia", "Lluvia 7 días (l/m²)", "Lluvia 30 días (l/m²)"]


def resumen_diario(series):
    """
    Calcula el resumen de cada día. Devuelve una lista de filas:
      {"fecha": date, "horas_dia": 24,
       "est": {clave: {"valores": [8 valores en el orden de SUB], "incompleto": [8 booleanos]}}}
    """
    por_dia = {nombre: agrupar_por_dia(serie) for nombre, serie in series.items()}
    todos = [d for pd in por_dia.values() for d in pd]
    if not todos:
        return []
    d, ultimo = min(todos), max(todos)
    dias = []
    while d <= ultimo:
        dias.append(d)
        d += timedelta(days=1)

    memoria = {clave: [] for clave, _ in ESTACIONES}  # (lluvia del día, ¿completo?)
    filas = []
    for d in dias:
        esperadas = horas_del_dia(d)
        fila = {"fecha": d, "horas_dia": esperadas, "est": {}}
        for clave, _ in ESTACIONES:
            temp = [v for v, c in por_dia[f"{clave}_temperatura"].get(d, []) if c == "ok"]
            lluvia = [v for v, c in por_dia[f"{clave}_lluvia"].get(d, []) if c == "ok"]
            temp_completa = len(temp) >= esperadas
            lluvia_completa = len(lluvia) >= esperadas

            valores = [
                min(temp) if temp else None,
                max(temp) if temp else None,
                sum(temp) / len(temp) if temp else None,
                len(temp),
                sum(lluvia) if lluvia else None,
                len(lluvia),
            ]
            memoria[clave].append((sum(lluvia) if lluvia else 0.0, lluvia_completa))
            v7, v30 = memoria[clave][-7:], memoria[clave][-30:]
            acum7 = sum(x for x, _ in v7) if len(v7) == 7 else None
            acum30 = sum(x for x, _ in v30) if len(v30) == 30 else None
            ok7 = len(v7) == 7 and all(c for _, c in v7)
            ok30 = len(v30) == 30 and all(c for _, c in v30)
            valores += [acum7, acum30]
            incompleto = [
                not temp_completa, not temp_completa, not temp_completa, not temp_completa,
                not lluvia_completa, not lluvia_completa,
                acum7 is not None and not ok7,
                acum30 is not None and not ok30,
            ]
            fila["est"][clave] = {"valores": valores, "incompleto": incompleto}
        filas.append(fila)
    return filas


def escribir_diario_csv(filas, ruta):
    """CSV del resumen diario, pensado para Power Query: coma como separador y punto decimal."""
    nombres = ["tmin", "tmax", "tmedia", "horas_t", "lluvia_dia", "horas_lluvia", "lluvia_7d", "lluvia_30d"]
    cab = ["fecha", "horas_dia"]
    for clave, _ in ESTACIONES:
        p = PREFIJOS[clave]
        cab += [f"{p}_{n}" for n in nombres] + [f"{p}_t_completa", f"{p}_lluvia_completa"]
    with ruta.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cab)
        for fila in filas:
            linea = [fila["fecha"].isoformat(), fila["horas_dia"]]
            for clave, _ in ESTACIONES:
                est = fila["est"][clave]
                for k, v in enumerate(est["valores"]):
                    if v is None:
                        linea.append("")
                    elif k in (3, 5):
                        linea.append(int(v))
                    else:
                        linea.append(f"{v:.2f}")
                linea += [0 if est["incompleto"][0] else 1, 0 if est["incompleto"][4] else 1]
            w.writerow(linea)


def hoja_diario(wb, filas):
    ws = wb.create_sheet("Diario")
    ws["A1"] = "Resumen diario (día local de España)"
    ws["A1"].font = F_TITULO
    ws["A2"] = (
        "Naranja = día con horas sin dato (resumen poco fiable). Los valores 'dudosos' de lluvia "
        "no se suman. El día de hoy siempre está incompleto. Ver hoja Leeme."
    )
    ws["A2"].font = Font(name=FUENTE, size=9, italic=True)

    ws.cell(row=4, column=1, value="Fecha")
    ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    col = 2
    for _, etiqueta in ESTACIONES:
        ws.cell(row=4, column=col, value=etiqueta)
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + len(SUB) - 1)
        for k, texto in enumerate(SUB):
            ws.cell(row=5, column=col + k, value=texto)
        col += len(SUB)
    ultima_col = col - 1
    for fila in (4, 5):
        for c in range(1, ultima_col + 1):
            celda = ws.cell(row=fila, column=c)
            celda.font, celda.alignment = F_CAB, CENTRO
            celda.fill = REL_GRUPO if fila == 4 else REL_CAB
    ws.row_dimensions[5].height = 42

    r = 6
    for fila in filas:
        c0 = ws.cell(row=r, column=1, value=fila["fecha"])
        c0.number_format, c0.font = "dd/mm/yyyy", F_NORMAL
        col = 2
        for clave, _ in ESTACIONES:
            est = fila["est"][clave]
            for k, valor in enumerate(est["valores"]):
                celda = ws.cell(row=r, column=col + k, value=valor)
                celda.font = F_NORMAL
                celda.number_format = "0" if k in (3, 5) else "0.0"
                if est["incompleto"][k]:
                    celda.fill = REL_INCOMPLETO
            col += len(SUB)
        r += 1

    ws.column_dimensions["A"].width = 12
    for c in range(2, ultima_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = "B6"
    return len(filas)


def hoja_leeme(wb, n_horas, n_dias):
    ws = wb.create_sheet("Leeme", 0)
    estado = {}
    ruta_estado = CARPETA / "estado.json"
    if ruta_estado.exists():
        estado = json.loads(ruta_estado.read_text(encoding="utf-8"))

    lineas = [
        ("Datos meteorológicos del Valle del Eria (SAIH Duero)", F_TITULO),
        ("", F_NORMAL),
        (f"Excel generado: {datetime.now(TZ).strftime('%d/%m/%Y %H:%M')} (hora de España)", F_NORMAL),
        (f"Última captura de datos (UTC): {estado.get('ejecucion_utc', 'sin datos')}", F_NORMAL),
        (f"Horas con algún dato en la hoja Horario: {n_horas} · Días en la hoja Diario: {n_dias}", F_NORMAL),
        ("", F_NORMAL),
        ("QUÉ HAY AQUÍ", F_NEGRITA),
        ("Diario: resumen por día de cada estación. Es la hoja para comparar con tus propias observaciones.", F_NORMAL),
        ("Horario: los datos hora a hora tal como los publica el SAIH. Las horas sin dato están VACÍAS (no son ceros).", F_NORMAL),
        ("Aparte, en el repositorio, datos/diario.csv tiene el mismo resumen diario y se actualiza cada día (para Power Query).", F_NORMAL),
        ("", F_NORMAL),
        ("ESTACIONES", F_NEGRITA),
        ("Castrocalbón (PL284): pluviómetro. Temperatura y lluvia.", F_NORMAL),
        ("Morales del Rey (EA082): estación de aforo del río Eria. Temperatura (sensor junto al río) y lluvia.", F_NORMAL),
        ("Ninguna de las dos está en tu punto de interés. Úsalas como tendencia, no como el dato exacto de cada lugar.", F_NORMAL),
        ("", F_NORMAL),
        ("COLORES", F_NEGRITA),
        ("Naranja claro: ese día tiene horas sin dato; el resumen puede no ser fiable.", F_NORMAL),
        ("Amarillo (hoja Horario): dato 'dudoso'. No se suma en los totales de lluvia.", F_NORMAL),
        ("", F_NORMAL),
        ("AVISOS IMPORTANTES", F_NEGRITA),
        ("Son datos en tiempo real, PROVISIONALES: el propio SAIH avisa de que pueden revisarse o corregirse.", F_NORMAL),
        ("Para estudios técnicos, el SAIH indica que los datos meteorológicos válidos son los de AEMET.", F_NORMAL),
        ("Los huecos suelen coincidir en las dos estaciones (fallos de comunicación), así que una no rellena a la otra.", F_NORMAL),
        ("Criterios míos (supuestos, ajustables en captura_saih.py): lluvia > 50 l/m² en una hora = 'dudoso'; "
         "temperatura fuera de -30..50 ºC = 'dudoso'.", F_NORMAL),
        ("El primer dato de lluvia de las dos estaciones (20/07/2026 00:00 = 120 l/m²) coincide en ambas y parece un "
         "artefacto: está marcado como dudoso.", F_NORMAL),
        ("Los totales de lluvia y las medias usan solo datos 'ok'. Un día sin todas sus horas sale en naranja.", F_NORMAL),
        ("", F_NORMAL),
        ("Este archivo se regenera entero: no escribas en él. Tus anotaciones van en otro archivo, aparte.", F_NEGRITA),
    ]
    if estado.get("problemas"):
        lineas += [("", F_NORMAL), ("PROBLEMAS DE LA ÚLTIMA CAPTURA", F_NEGRITA)]
        lineas += [(f"- {p}", F_NORMAL) for p in estado["problemas"]]
    for i, (texto, fuente) in enumerate(lineas, start=1):
        celda = ws.cell(row=i, column=1, value=texto)
        celda.font = fuente
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120


def main():
    series = cargar()
    if not any(series.values()):
        print("No hay datos en la carpeta de datos. Ejecuta antes captura_saih.py.")
        return 1

    # El resumen diario en CSV se actualiza SIEMPRE (es lo que lee Power Query).
    filas_diario = resumen_diario(series)
    escribir_diario_csv(filas_diario, CARPETA / "diario.csv")
    print(f"CSV diario guardado: {CARPETA / 'diario.csv'} ({len(filas_diario)} días)")

    # Para no engordar el repositorio, el Excel se rehace solo los lunes,
    # salvo que se fuerce (ejecución manual) o que todavía no exista.
    forzar = os.environ.get("FORZAR_EXCEL", "0") == "1"
    es_lunes = datetime.now(TZ).weekday() == 0
    if not (forzar or es_lunes or not SALIDA.exists()):
        print("Hoy no toca rehacer el Excel (solo los lunes o al forzarlo a mano). Se omite.")
        return 0

    wb = Workbook()
    wb.remove(wb.active)
    n_horas = hoja_horario(wb, series)
    n_dias = hoja_diario(wb, filas_diario)
    hoja_leeme(wb, n_horas, n_dias)
    wb.move_sheet("Diario", offset=-(wb.sheetnames.index("Diario") - 1))  # Leeme, Diario, Horario
    wb.save(SALIDA)
    print(f"Excel guardado: {SALIDA} ({n_horas} horas, {n_dias} días)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
